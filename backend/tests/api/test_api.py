"""API: чтение без входа, режим админа, аудит, экспорт, демо-эмулятор."""

from __future__ import annotations

import io

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from teamtime.services.ingest import ingest
from tests import factories as f
from tests.api.conftest import ADMIN_HEADERS

pytestmark = pytest.mark.db


async def seed(maker: async_sessionmaker[AsyncSession]) -> int:
    async with maker() as s:
        await f.forum(s)
        e = await f.employee(s, "Иванов И.И.", tg_id=501)
        await ingest(s, f.message(f.schedule_text(), thread=f.SCHEDULE_THREAD))
        await ingest(s, f.message(f.daily_text("26.10.2026", "10:15 - 18:00", 80)))
        await s.commit()
        return e.id


async def test_health_and_openapi(client: httpx.AsyncClient) -> None:
    r = await client.get("/health")
    assert r.json()["db"] == "ok"
    spec = (await client.get("/openapi.json")).json()
    assert spec["info"]["title"] == "TeamTime API"
    assert "/api/v1/weeks/{iso_year}/{iso_week}" in spec["paths"]


async def test_read_screens_without_login(client: httpx.AsyncClient, maker) -> None:
    emp_id = await seed(maker)
    meta = (await client.get("/meta")).json()
    assert meta["is_admin"] is False
    assert meta["today"] == "2026-10-26"

    week = (await client.get("/weeks/2026/44")).json()
    assert week["label"] == "44 неделя"
    cell = week["rows"][0]["cells"][0]
    assert cell["late_sec"] == 15 * 60
    assert cell["fact"]["result_sec"] == int(7.75 * 3600 * 0.8)

    assert (await client.get("/months/2026/10")).status_code == 200
    assert (await client.get("/today")).json()["counts"]["dailies"] == 1
    assert (await client.get(f"/employees/{emp_id}")).json()["employee"][
        "short_name"
    ] == "Иванов И.И."
    assert len((await client.get("/plans/2026/44")).json()["rows"][0]["versions"]) == 1
    cal = (
        await client.get("/calendar", params={"start": "2026-10-26", "end": "2026-11-01"})
    ).json()
    assert len(cal[str(emp_id)]) == 7
    raw_id = cell["source"]["daily_raw_message_id"]
    msg = (await client.get(f"/messages/{raw_id}")).json()
    assert msg["versions"][0]["text"].startswith("26.10.2026")
    hits = (await client.get("/search", params={"q": "Иванов"})).json()["hits"]
    assert hits[0]["href"] == f"/people/{emp_id}"
    assert (await client.get("/penalty", params={"late_minutes": 340})).json()[
        "penalty_sec"
    ] == 20742


async def test_bad_week_is_422(client: httpx.AsyncClient) -> None:
    assert (await client.get("/weeks/2026/54")).status_code == 422


async def test_writes_need_admin_and_csrf_header(client: httpx.AsyncClient, maker) -> None:
    emp_id = await seed(maker)
    body = {"kind": "sick", "note": ""}
    url = f"/admin/days/{emp_id}/2026-10-27/status"
    assert (await client.put(url, json=body, headers=ADMIN_HEADERS)).status_code == 401
    assert (await client.post("/auth/login", json={"password": "wrong"})).status_code == 401
    assert (await client.post("/auth/login", json={"password": "secret"})).status_code == 200
    assert (await client.put(url, json=body)).status_code == 403  # нет X-Requested-With
    assert (await client.put(url, json=body, headers=ADMIN_HEADERS)).status_code == 200
    audit = (await client.get("/audit")).json()
    assert audit[0]["action"] == "status.set"
    assert audit[0]["actor"] == "admin"


async def test_login_throttling(client: httpx.AsyncClient) -> None:
    codes = [
        (await client.post("/auth/login", json={"password": "x"})).status_code for _ in range(6)
    ]
    assert codes[-1] == 429


async def test_admin_employee_crud_and_validation(admin: httpx.AsyncClient) -> None:
    body = {
        "full_name": "Сидорова Анна Петровна",
        "short_name": "Сидорова А.П.",
        "position": "Дизайнер",
        "monthly_norm_minutes": 4800,
        "hourly_rate": "700",
    }
    r = await admin.post("/admin/employees", json=body, headers=ADMIN_HEADERS)
    assert r.status_code == 201, r.text
    assert r.json()["match_key"] == "сидорова ап"
    dup = await admin.post("/admin/employees", json=body, headers=ADMIN_HEADERS)
    assert dup.status_code == 422
    bad = await admin.post(
        "/admin/employees", json={**body, "short_name": "Анна"}, headers=ADMIN_HEADERS
    )
    assert bad.status_code == 422


async def test_admin_fix_fact_and_settings(admin: httpx.AsyncClient, maker) -> None:
    emp_id = await seed(maker)
    r = await admin.put(
        f"/admin/days/{emp_id}/2026-10-26/fact",
        json={"intervals": [[36000, 64800]], "keep_efficiency": True, "reason": "проверка"},
        headers=ADMIN_HEADERS,
    )
    assert r.status_code == 200, r.text
    week = (await admin.get("/weeks/2026/44")).json()
    assert week["rows"][0]["cells"][0]["late_sec"] == 0

    s = (await admin.get("/admin/settings")).json()
    s["late_tolerance_min"] = 5
    s["penalty"]["quad_coef"] = "1/7"
    r = await admin.put("/admin/settings", json=s, headers=ADMIN_HEADERS)
    assert r.status_code == 200
    assert r.json()["warnings"]  # шкала стала разрывной — предупреждаем
    assert (await admin.get("/meta")).json()["late_tolerance_min"] == 5


async def test_exports(client: httpx.AsyncClient, maker) -> None:
    await seed(maker)
    csv = await client.get("/export/weeks/2026/44.csv")
    assert csv.headers["content-type"].startswith("text/csv")
    text = csv.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Сотрудник;Должность;Дата")
    assert "Иванов И.И.;Разработчик;26.10.2026;Пн;10:00–18:00;10:15;18:00" in text
    xlsx = await client.get("/export/months/2026/10.xlsx")
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Итог за месяц", "Опоздания", "Дни"]
    week_xlsx = load_workbook(io.BytesIO((await client.get("/export/weeks/2026/44.xlsx")).content))
    ws = week_xlsx["44 неделя"]
    assert ws["A2"].value == "Иванов И.И."
    assert ws["F2"].number_format == "[h]:mm"


async def test_demo_emulator(client: httpx.AsyncClient, maker) -> None:
    await seed(maker)
    r = await client.post(
        "/demo/messages",
        json={
            "telegram_user_id": 501,
            "topic": "daily",
            "text": f.daily_text("23.10.2026"),
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["reaction"] == "👌"
    r = await client.post(
        "/demo/messages",
        json={
            "telegram_user_id": 501,
            "topic": "daily",
            "text": "23.10.2026\nбез времени",
        },
    )
    body = r.json()
    assert body["reaction"] == "🤔"
    assert "Исправить в боте" in body["reply"]


async def test_inbox_flow(admin: httpx.AsyncClient, maker) -> None:
    async with maker() as s:
        await f.forum(s)
        await f.employee(s, tg_id=501)
        await ingest(s, f.message("23.10.2026\n10:00 - 18:00\n- без эффективности"))
        await s.commit()
    inbox = (await admin.get("/inbox")).json()
    assert inbox["counts"]["pending"] == 1
    item = inbox["items"][0]
    assert item["parse"]["status"] == "partial"
    r = await admin.post(
        f"/admin/inbox/{item['parse']['id']}/resolve",
        json={"resolution": "accepted"},
        headers=ADMIN_HEADERS,
    )
    assert r.status_code == 200
    assert (await admin.get("/inbox")).json()["counts"]["pending"] == 0


async def test_import_upload(admin: httpx.AsyncClient, maker) -> None:
    from pathlib import Path

    async with maker() as s:
        await f.employee(s, "Иванов И.И.", tg_id=501)
        await s.commit()
    data = (Path(__file__).resolve().parents[1] / "fixtures" / "telegram_export.json").read_bytes()
    files = {"file": ("result.json", data, "application/json")}
    r = await admin.post("/admin/import/telegram", files=files, headers=ADMIN_HEADERS)
    assert r.status_code == 200, r.text
    assert r.json()["imported"] == 4
    bad = await admin.post(
        "/admin/import/telegram",
        files={"file": ("x.json", b"{}", "application/json")},
        headers=ADMIN_HEADERS,
    )
    assert bad.status_code == 422
    assert (await admin.get("/audit")).json()[0]["action"] == "import.telegram"
