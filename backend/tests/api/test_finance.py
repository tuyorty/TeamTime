"""Расчёт оплаты, история условий и изменения премий на вымышленных данных."""

import asyncio
import io
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from teamtime.db.models import AuditLog, EmployeeTerms, ProjectBonus
from teamtime.services.ingest import ingest
from tests import factories as f
from tests.api.conftest import ADMIN_HEADERS

pytestmark = pytest.mark.db


def bonus_payload(employee_id, **updates):
    return {
        "employee_id": employee_id,
        "request_id": str(uuid4()),
        "month": "2026-10-01",
        "title": "Демо-проект",
        "amount": "1000.25",
        "note": "Запуск модуля",
        **updates,
    }


async def test_full_salary_historical_rate_and_overtime_match_month(client, maker):
    async with maker() as session:
        await f.forum(session)
        employee = await f.employee(session, norm_hours=160, rate=900, tg_id=501)
        session.add(
            EmployeeTerms(
                employee_id=employee.id,
                effective_from=date(2000, 1, 1),
                monthly_norm_minutes=4 * 60,
                hourly_rate=Decimal("500.01"),
            )
        )
        session.add(
            EmployeeTerms(
                employee_id=employee.id,
                effective_from=date(2026, 11, 1),
                monthly_norm_minutes=160 * 60,
                hourly_rate=Decimal(900),
            )
        )
        await ingest(session, f.message(f.schedule_text(), thread=f.SCHEDULE_THREAD))
        await ingest(session, f.message(f.daily_text("26.10.2026", "10:20 - 18:00", 80)))
        await session.commit()
    finance = (await client.get("/finances/2026/10")).json()
    month = (await client.get("/months/2026/10?history=1")).json()
    row = finance["rows"][0]
    assert row["summary"] == month["rows"][0]["summary"]
    assert row["amounts"]["salary"] == "2000.04"
    assert row["amounts"]["overtime_pay"] == month["rows"][0]["summary"]["bonus"]
    assert Decimal(row["amounts"]["total"]) == Decimal("2000.04") + Decimal(
        row["amounts"]["overtime_pay"]
    )
    november = (await client.get("/finances/2026/11")).json()
    assert november["rows"][0]["amounts"] == {
        "salary": "144000.00",
        "overtime_pay": "0.00",
        "project_bonuses": "0.00",
        "total": "144000.00",
    }


async def test_bonus_retry_edit_conflict_cancel_and_audit(admin, maker):
    async with maker() as session:
        employee = await f.employee(session)
        await session.commit()
    payload = bonus_payload(employee.id)
    responses = await asyncio.gather(
        *(admin.post("/admin/bonuses", json=payload, headers=ADMIN_HEADERS) for _ in range(2))
    )
    assert [r.status_code for r in responses] == [200, 200]
    ids = {r.json()["id"] for r in responses}
    assert len(ids) == 1
    bonus_id = ids.pop()
    conflict = await admin.post(
        "/admin/bonuses", json={**payload, "amount": "2000"}, headers=ADMIN_HEADERS
    )
    assert conflict.status_code == 409
    edit = {"version": 1, "title": "Проект завершён", "amount": "1500.75", "note": "Уточнили сумму"}
    edited = await admin.put(f"/admin/bonuses/{bonus_id}", json=edit, headers=ADMIN_HEADERS)
    assert edited.status_code == 200
    assert edited.json()["version"] == 2
    assert (
        await admin.put(f"/admin/bonuses/{bonus_id}", json=edit, headers=ADMIN_HEADERS)
    ).status_code == 409
    report = (await admin.get("/finances/2026/10")).json()
    assert report["totals"]["total"] == "81500.75"
    cancelled = await admin.post(
        f"/admin/bonuses/{bonus_id}/cancel",
        json={"version": 2, "reason": "Назначили ошибочно"},
        headers=ADMIN_HEADERS,
    )
    assert cancelled.status_code == 200
    report = (await admin.get("/finances/2026/10")).json()
    assert report["totals"]["total"] == "80000.00"
    assert report["rows"][0]["bonuses"][0]["cancel_reason"] == "Назначили ошибочно"
    assert (
        await admin.post("/admin/bonuses", json=payload, headers=ADMIN_HEADERS)
    ).status_code == 409
    async with maker() as session:
        assert await session.scalar(select(func.count()).select_from(ProjectBonus)) == 1
        logs = list(
            await session.scalars(
                select(AuditLog)
                .where(AuditLog.entity_type == "project_bonus")
                .order_by(AuditLog.id)
            )
        )
        assert [log.action for log in logs] == ["bonus.created", "bonus.updated", "bonus.cancelled"]
        assert logs[1].before["amount"] == "1000.25"
        assert logs[1].after["amount"] == "1500.75"


async def test_finance_read_is_public_but_bonus_write_requires_admin_and_csrf(client):
    assert (await client.get("/finances/2026/10")).status_code == 200
    payload = bonus_payload(1)
    for method, url, body in [
        ("post", "/admin/bonuses", payload),
        ("put", "/admin/bonuses/1", {**payload, "version": 1}),
        ("post", "/admin/bonuses/1/cancel", {"version": 1, "reason": "Ошибка"}),
    ]:
        assert (
            await getattr(client, method)(url, json=body, headers=ADMIN_HEADERS)
        ).status_code == 401
    await client.post("/auth/login", json={"password": "secret"})
    assert (await client.post("/admin/bonuses", json=payload)).status_code == 403


@pytest.mark.parametrize(
    "updates",
    [
        {"amount": "0"},
        {"amount": "-1"},
        {"amount": "1.001"},
        {"amount": "NaN"},
        {"title": "   "},
        {"month": "2026-10-15"},
        {"month": "1999-12-01"},
    ],
)
async def test_invalid_bonus_is_rejected(admin, updates):
    response = await admin.post(
        "/admin/bonuses", json=bonus_payload(1, **updates), headers=ADMIN_HEADERS
    )
    assert response.status_code == 422


async def test_archived_employee_bonus_remains_visible_and_export_is_filtered(admin, maker):
    async with maker() as session:
        employee = await f.employee(session, "Тестов Т.Т.", norm_hours=0, rate=0)
        employee.is_active = False
        await f.employee(session, "Примеров П.П.")
        await session.commit()
    await admin.post(
        "/admin/bonuses", json=bonus_payload(employee.id, title="=1+1"), headers=ADMIN_HEADERS
    )
    report = (await admin.get("/finances/2026/10")).json()
    assert len(report["rows"]) == 2
    row = next(r for r in report["rows"] if r["employee"]["id"] == employee.id)
    assert row["amounts"]["total"] == "1000.25"
    previous = (await admin.get("/finances/2026/9")).json()
    assert len(previous["rows"]) == 1
    response = await admin.get(f"/export/finances/2026/10.xlsx?employee_id={employee.id}")
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert wb["Расчёт оплаты"].max_row == 3
    assert wb["Расчёт оплаты"].cell(3, 9).value == 1000.25
    assert wb["Проектные премии"].cell(2, 3).data_type == "s"
    csv = await admin.get(f"/export/finances/2026/10.csv?employee_id={employee.id}")
    assert "1000.25" in csv.content.decode("utf-8-sig")
