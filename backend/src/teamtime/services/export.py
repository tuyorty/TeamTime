"""Экспорт недели и месяца в CSV (UTF-8 с BOM, «;» — открывается в Excel) и XLSX.

В XLSX длительности — настоящие значения времени Excel в формате ``[h]:mm``: их можно
складывать формулами, а Excel, как и исходная таблица, не округляет секунды до минут.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from teamtime.calc import STATUS_LABELS
from teamtime.calc.duration import format_hm
from teamtime.schemas import DayCellOut, MonthOut, WeekOut

WEEKDAYS = ("Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс")
DURATION_FMT = "[h]:mm"
TIME_FMT = "h:mm"
HEADER_FILL = PatternFill("solid", fgColor="E8EAEE")
HEADER_FONT = Font(bold=True)
THIN = Side(style="thin", color="C9CDD4")


class Duration(int):
    """Длительность в секундах: в CSV — «Ч:ММ», в XLSX — значение времени."""


class TimeOfDay(int):
    """Время суток в секундах от полуночи."""


def _plan_text(c: DayCellOut) -> str:
    if c.plan is None:
        return ""
    if c.plan.kind.value == "off":
        return "выходной"
    if c.plan.kind.value == "note":
        return c.plan.note or ""
    return f"{format_hm(c.plan.start)}–{format_hm(c.plan.end)}"


def day_rows(name: str, position: str, cells: Sequence[DayCellOut]) -> list[list[Any]]:
    rows = []
    for c in cells:
        f = c.fact
        rows.append([
            name,
            position,
            c.date,
            WEEKDAYS[c.date.weekday()],
            _plan_text(c),
            TimeOfDay(f.start) if f else None,
            TimeOfDay(f.end) if f else None,
            "; ".join(f"{format_hm(a)}–{format_hm(b)}" for a, b in f.intervals) if f else "",
            Duration(f.worked_sec) if f else None,
            f.efficiency if f else None,
            Duration(f.result_sec) if f and f.result_sec is not None else None,
            Duration(c.late_sec) if c.late_sec else None,
            STATUS_LABELS[c.status] if c.status else "",
            " · ".join(f.tasks) if f else "",
        ])  # fmt: skip
    return rows


DAY_HEADER = [
    "Сотрудник", "Должность", "Дата", "День", "План", "Пришёл", "Ушёл", "Интервалы", "Часы",
    "e%", "Итог", "Опоздание", "Статус", "Сделано",
]  # fmt: skip


def week_day_rows(week: WeekOut) -> list[list[Any]]:
    return [
        row
        for r in week.rows
        for row in day_rows(r.employee.short_name, r.employee.position, r.cells)
    ]


def month_summary_rows(month: MonthOut) -> tuple[list[str], list[list[Any]]]:
    header = [
        "Сотрудник", "Должность", "Часы(+)", "Часы(e%)", "Опоздания(−)", "Итого", "Норма",
        "Выполнение нормы", "Опозданий, дней", "Сумма опозданий", "Переработка",
        "Ставка, ₽/ч", "Премия, ₽",
    ]  # fmt: skip
    rows = [
        [
            r.employee.short_name,
            r.employee.position,
            Duration(r.summary.worked_sec),
            Duration(r.summary.result_sec),
            Duration(r.summary.penalty_sec),
            Duration(r.summary.total_sec),
            Duration(r.summary.norm_sec),
            r.summary.norm_progress,
            r.summary.late_count,
            Duration(r.summary.late_total_sec),
            Duration(r.summary.overtime_sec),
            float(r.summary.hourly_rate),
            float(r.summary.bonus),
        ]
        for r in month.rows
    ]
    return header, rows


def month_late_rows(month: MonthOut) -> tuple[list[str], list[list[Any]]]:
    header = ["Сотрудник", "Дата", "План", "Пришёл", "Опоздание"]
    rows = [
        [
            r.employee.short_name,
            d.date,
            TimeOfDay(d.plan_start),
            TimeOfDay(d.fact_start),
            Duration(d.late_sec),
        ]
        for r in month.rows
        for d in r.late_days
    ]
    return header, rows


# --- форматы -------------------------------------------------------------------------------


# Текст из Telegram (задачи, заметки) пишет любой участник чата. Ячейку, которая начинается
# с этих символов, Excel и LibreOffice выполняют как формулу (=HYPERLINK, =WEBSERVICE…).
FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n")


def _defuse(text: str) -> str:
    """Текст, который табличный редактор не примет за формулу (рекомендация OWASP)."""
    return f"'{text}" if text.startswith(FORMULA_START) else text


def _csv_value(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return _defuse(v)
    if isinstance(v, Duration | TimeOfDay):
        return format_hm(int(v))
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float):
        return f"{v:.2f}".replace(".", ",")
    return str(v)


def to_csv(header: Sequence[str], rows: Sequence[Sequence[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerow(header)
    for row in rows:
        writer.writerow([_csv_value(v) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _write_sheet(ws: Any, header: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    ws.append(list(header))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    for row in rows:
        out: list[Any] = []
        for v in row:
            out.append(int(v) / 86400 if isinstance(v, Duration | TimeOfDay) else v)
        ws.append(out)
        r = ws.max_row
        for col, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=col)
            if isinstance(v, str) and cell.data_type == "f":
                cell.data_type = "s"  # openpyxl принял «=…» за формулу: храним как текст
            elif isinstance(v, Duration):
                cell.number_format = DURATION_FMT
            elif isinstance(v, TimeOfDay):
                cell.number_format = TIME_FMT
            elif isinstance(v, date):
                cell.number_format = "DD.MM.YYYY"
            elif header[col - 1] == "Выполнение нормы":
                cell.number_format = "0%"
            elif header[col - 1].endswith("₽") or header[col - 1].startswith("Ставка"):
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "B2"
    for col in range(1, len(header) + 1):
        values = [
            str(ws.cell(row=r, column=col).value or "") for r in range(1, min(ws.max_row, 60) + 1)
        ]
        ws.column_dimensions[get_column_letter(col)].width = min(
            48, max(9, *(len(v) for v in values)) + 2
        )


def to_xlsx(sheets: Sequence[tuple[str, Sequence[str], Sequence[Sequence[Any]]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        _write_sheet(wb.create_sheet(title[:31]), header, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def week_grid_rows(week: WeekOut) -> tuple[list[str], list[list[Any]]]:
    """Лист «как в таблице»: по строке на сотрудника, по блоку колонок на день."""
    header = ["Сотрудник", "Должность"]
    for d in week.days:
        prefix = f"{WEEKDAYS[d.weekday()]} {d:%d.%m}"
        header += [f"{prefix} план", f"{prefix} пришёл", f"{prefix} ушёл", f"{prefix} часы",
                   f"{prefix} e%", f"{prefix} итог"]  # fmt: skip
    header += ["Часы за неделю", "Итог за неделю", "Опоздания"]
    rows = []
    for r in week.rows:
        row: list[Any] = [r.employee.short_name, r.employee.position]
        for c in r.cells:
            f = c.fact
            row += [
                STATUS_LABELS[c.status] if c.status and not f else _plan_text(c),
                TimeOfDay(f.start) if f else None,
                TimeOfDay(f.end) if f else None,
                Duration(f.worked_sec) if f else None,
                f.efficiency if f else None,
                Duration(f.result_sec) if f and f.result_sec is not None else None,
            ]
        row += [
            Duration(r.totals.worked_sec),
            Duration(r.totals.result_sec),
            Duration(r.totals.late_sec),
        ]
        rows.append(row)
    return header, rows
